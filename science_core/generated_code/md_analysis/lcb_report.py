#!/usr/bin/env python3
"""Generate a lightweight project report from an LCB integrated output."""

from __future__ import annotations

import argparse
import csv
import json
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from lcb_project_config import load_lcb_config


SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parents[1]


def run(command: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(command, text=True, capture_output=True, cwd=ROOT)


def read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, encoding="utf-8-sig")


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def present(value: Any) -> bool:
    text = str(value if value is not None else "").strip()
    return text != "" and text.lower() not in {"nan", "none", "missing"}


def figure_rows(figures_dir: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not figures_dir.exists():
        return rows
    for path in sorted(figures_dir.rglob("*")):
        if path.suffix.lower() not in {".png", ".pdf"}:
            continue
        rows.append(
            {
                "figure": path.name,
                "relative_path": str(path.relative_to(figures_dir.parent)).replace("\\", "/"),
                "format": path.suffix.lower().lstrip("."),
                "size_bytes": path.stat().st_size,
            }
        )
    return rows


def status_summary(wide: pd.DataFrame) -> tuple[list[str], dict[str, Counter]]:
    excluded = {"status_ok_count", "status_warn_count", "status_fail_count", "status_missing_count"}
    status_cols = [col for col in wide.columns if col.startswith("status_") and col not in excluded]
    counts = {col: Counter(wide[col].fillna("MISSING").astype(str)) for col in status_cols}
    lines: list[str] = []
    for col in sorted(status_cols):
        pieces = ", ".join(f"{key}={value}" for key, value in sorted(counts[col].items()))
        lines.append(f"- `{col}`: {pieces}")
    return lines, counts


def make_markdown(
    config,
    wide: pd.DataFrame,
    figures: list[dict[str, Any]],
    aggregate_dir: Path,
    source_root: Path,
    release_manifest: dict[str, Any] | None = None,
) -> str:
    generated = datetime.now().isoformat(timespec="seconds")
    total = int(len(wide))
    analyzed = int(wide.get("any_analyzed", pd.Series([False] * total)).fillna(False).astype(bool).sum()) if total else 0
    status_lines, _ = status_summary(wide)
    quality_counts = Counter()
    if "quality_flags" in wide.columns:
        for value in wide["quality_flags"].fillna("").astype(str):
            for flag in [part for part in value.split(";") if part]:
                quality_counts[flag] += 1

    lines = [
        f"# {config.project_id.upper()} Project Report",
        "",
        f"- Generated: `{generated}`",
        f"- Manifest: `{config.manifest_csv}`",
        f"- Source root: `{source_root}`",
        f"- Total design cases: `{total}`",
        f"- Cases with any analyzed observable: `{analyzed}`",
        "",
        "## Observable Status",
        "",
        *status_lines,
        "",
        "## Quality Flags",
        "",
    ]
    if quality_counts:
        for flag, count in sorted(quality_counts.items()):
            lines.append(f"- `{flag}`: {count}")
    else:
        lines.append("- No quality flags found in the current summary table.")
    missing_cols = [col for col in ["status_Rg", "status_RDF", "status_MSD", "status_shape", "status_dynamics", "status_contact", "status_density"] if col in wide.columns]
    if missing_cols:
        missing = wide[wide[missing_cols].apply(lambda row: any(str(v).upper() in {"MISSING", "FAIL", ""} for v in row), axis=1)]
        lines.extend(["", "## Cases With Missing Or Failed Non-Viscosity Outputs", ""])
        if len(missing):
            for _, row in missing.head(40).iterrows():
                bad = [col.replace("status_", "") for col in missing_cols if str(row.get(col, "")).upper() in {"MISSING", "FAIL", ""}]
                lines.append(f"- `{row.get('case_id')}`: {','.join(bad)}")
            if len(missing) > 40:
                lines.append(f"- ... {len(missing) - 40} more")
        else:
            lines.append("- None.")
    provenance = read_csv(aggregate_dir / "provenance_cases.csv")
    metrics = read_csv(aggregate_dir / "automation_metrics.csv")
    ml_mask = read_csv(aggregate_dir / "ml_quality_mask.csv")
    lines.extend(["", "## Automation And Provenance", ""])
    if not provenance.empty:
        pub_count = int(provenance.get("publication_grade", pd.Series(dtype=bool)).astype(str).str.lower().isin(["true", "1"]).sum())
        ml_count = int(provenance.get("ml_usable", pd.Series(dtype=bool)).astype(str).str.lower().isin(["true", "1"]).sum())
        lines.append(f"- Provenance ledger: `{aggregate_dir / 'provenance_cases.csv'}`")
        lines.append(f"- Publication-grade structural cases: `{pub_count}`")
        lines.append(f"- ML-usable cases by default mask: `{ml_count}`")
        lines.append(f"- Provenance graph: `{aggregate_dir / 'provenance_nodes.csv'}`, `{aggregate_dir / 'provenance_edges.csv'}`")
    else:
        lines.append("- Provenance ledger not found. Run `lcb_publication_assets.py` or rebuild overview.")
    if not metrics.empty:
        for name in ["total_design_cases", "any_analyzed_cases", "strict_done_cases", "active_processes", "ready_for_analysis_cases"]:
            sub = metrics[metrics["metric"] == name]
            if not sub.empty:
                lines.append(f"- `{name}`: {sub.iloc[0].get('value')}")
        lines.append(f"- Automation metrics: `{aggregate_dir / 'automation_metrics.md'}`")
    if not ml_mask.empty:
        lines.append(f"- ML dataset: `{aggregate_dir / 'ml_features.csv'}`, `{aggregate_dir / 'ml_targets.csv'}`, `{aggregate_dir / 'ml_quality_mask.csv'}`")
    if release_manifest:
        lines.extend(
            [
                "",
                "## Release Freeze",
                "",
                f"- Release level: `{release_manifest.get('release_level', '')}`",
                f"- Freeze root: `{release_manifest.get('freeze_root', '')}`",
                f"- Artifact count: `{release_manifest.get('artifact_count', '')}`",
            ]
        )
    if "rdf_validation_status" in wide.columns:
        counts = Counter(wide["rdf_validation_status"].fillna("").astype(str))
        pieces = ", ".join(f"{key}={value}" for key, value in sorted(counts.items()))
        lines.extend(["", "## RDF Review Status", "", f"- `{pieces}`"])
    lines.extend(["", "## Figures", ""])
    if figures:
        for row in figures[:80]:
            lines.append(f"- `{row['relative_path']}`")
        if len(figures) > 80:
            lines.append(f"- ... {len(figures) - 80} more")
    else:
        lines.append("- No PNG/PDF figures found under aggregate/plots.")
    lines.extend(
        [
            "",
            "## Notes",
            "",
            "- RDF is included in the automated postprocessing package; `rdf_validation_status` records current/validated/excluded audit state.",
            "- Viscosity is treated as historical/exploratory unless explicitly regenerated.",
            "- Entanglement and packing proxy outputs are not part of the default completion criteria.",
            "- The report is generated from existing integrated outputs; it does not submit jobs or modify remote state.",
        ]
    )
    return "\n".join(lines) + "\n"


def maybe_build_pptx(config, report_md: Path, figures: list[dict[str, Any]], figures_root: Path, out_path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.util import Inches, Pt  # type: ignore
    except Exception as exc:  # noqa: BLE001
        return f"python-pptx unavailable: {exc}"

    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = f"{config.project_id.upper()} Report"
    title.placeholders[1].text = "Generated from integrated LCB outputs"

    summary_slide = prs.slides.add_slide(prs.slide_layouts[1])
    summary_slide.shapes.title.text = "Report Files"
    body = summary_slide.placeholders[1].text_frame
    body.clear()
    for text in [f"Markdown: {report_md.name}", f"Integrated: {config.integrated_root}", "Viscosity: historical/exploratory"]:
        para = body.add_paragraph()
        para.text = text
        para.font.size = Pt(18)

    for row in figures[:12]:
        if row["format"] != "png":
            continue
        path = figures_root.parent / row["relative_path"]
        if not path.exists():
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = path.stem[:60]
        slide.shapes.add_picture(str(path), Inches(0.5), Inches(1.0), width=Inches(9.0))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return f"wrote {out_path}"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=SCRIPT_DIR / "configs" / "lcb300.json")
    parser.add_argument("--rebuild-overview", action="store_true", help="Run build_lcb_overview.py before reporting.")
    parser.add_argument("--pptx", action="store_true", help="Also generate a lightweight PPTX if python-pptx is installed.")
    parser.add_argument("--release", type=Path, default=None, help="Read/write report files inside a freeze directory.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    config = load_lcb_config(args.config)
    if args.rebuild_overview:
        proc = run([sys.executable, str(SCRIPT_DIR / "build_lcb_overview.py"), "--config", str(args.config)])
        if proc.returncode != 0:
            print(proc.stdout)
            print(proc.stderr, file=sys.stderr)
            return proc.returncode

    release_manifest: dict[str, Any] | None = None
    if args.release:
        release_root = args.release.resolve()
        aggregate_dir = release_root / "aggregate"
        summary_path = aggregate_dir / f"{config.project_id}_summary_wide.csv"
        figures_dir = release_root / "figures"
        if not figures_dir.exists():
            figures_dir = aggregate_dir / "plots"
        report_root = release_root
        manifest_path = release_root / "freeze_manifest.json"
        if manifest_path.exists():
            release_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    else:
        release_root = config.integrated_root
        aggregate_dir = config.integrated_root / "aggregate"
        summary_path = config.integrated_summary_csv
        figures_dir = aggregate_dir / "plots"
        report_root = aggregate_dir

    wide = read_csv(summary_path)
    if wide.empty:
        print(f"Integrated summary is missing or empty: {summary_path}", file=sys.stderr)
        return 2
    figures = figure_rows(figures_dir)
    report_path = report_root / "project_report.md"
    report_path.write_text(make_markdown(config, wide, figures, aggregate_dir, release_root, release_manifest), encoding="utf-8")
    write_csv(report_root / "figure_index.csv", figures, ["figure", "relative_path", "format", "size_bytes"])
    manifest = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "project_id": config.project_id,
        "report_md": str(report_path),
        "figure_count": len(figures),
        "summary_csv": str(summary_path),
        "release_root": str(args.release.resolve()) if args.release else "",
    }
    if args.pptx:
        pptx_path = report_root / f"{config.project_id}_report.pptx"
        manifest["pptx_status"] = maybe_build_pptx(config, report_path, figures, figures_dir, pptx_path)
    (report_root / "project_report_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {report_path}")
    print(f"Wrote {report_root / 'figure_index.csv'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
